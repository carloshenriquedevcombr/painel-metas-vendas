from flask import Flask, render_template, request, jsonify, redirect, url_for, flash, send_file
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
import sqlite3
from datetime import datetime, timedelta
import pytz 
import io
import os
import secrets
from openpyxl import Workbook, load_workbook

# IMPORTAÇÃO DO NOSSO "CÉREBRO" DE BI
from core import calculos

app = Flask(__name__)
# Segredos vêm de variáveis de ambiente (veja .env.example)
app.secret_key = os.environ.get('SECRET_KEY')
if not app.secret_key:
    # Sem SECRET_KEY: gera uma chave aleatória (as sessões caem a cada reinício)
    app.secret_key = secrets.token_hex(32)
    print("[AVISO] SECRET_KEY não definida; usando chave temporária aleatória.")

# Senha usada ao criar usuário sem senha ou ao resetar senha pelo admin
SENHA_PADRAO = os.environ.get('SENHA_PADRAO')
if not SENHA_PADRAO:
    SENHA_PADRAO = secrets.token_urlsafe(9)
    print(f"[AVISO] SENHA_PADRAO não definida; senha padrão temporária desta execução: {SENHA_PADRAO}")

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'

# ==========================================
# ARQUITETURA DE BANCOS SEPARADOS (ISOLAMENTO)
# ==========================================
# Pasta dos bancos SQLite (criados vazios na primeira execução)
DATA_DIR = os.environ.get('DATA_DIR', os.path.dirname(os.path.abspath(__file__)))
os.makedirs(DATA_DIR, exist_ok=True)
DB_USERS = os.path.join(DATA_DIR, 'db_usuarios.db')
DB_SALES = os.path.join(DATA_DIR, 'db_vendas.db')
DB_METAS = os.path.join(DATA_DIR, 'db_metas.db')
DB_LOGS  = os.path.join(DATA_DIR, 'db_logs.db')

# Mapeamento para o sistema de Backups e Importação
DB_MAP = {
    'vendas': (DB_SALES, ['resultados']),
    'metas': (DB_METAS, ['metas', 'metas_semestrais_config', 'config_sistema']),
    'usuarios': (DB_USERS, ['usuarios']),
    'logs': (DB_LOGS, ['logs'])
}

def get_db(db_name):
    conn = sqlite3.connect(db_name, timeout=15)
    conn.row_factory = sqlite3.Row
    return conn

# --- HELPERS ---
def check_admin_access():
    return current_user.is_authenticated and current_user.role == 'admin'

def registrar_log(usuario_responsavel, acao, detalhes):
    try:
        conn = get_db(DB_LOGS)
        cursor = conn.cursor()
        tz = pytz.timezone('America/Sao_Paulo')
        data_hora = datetime.now(tz).strftime("%d/%m/%Y %H:%M") 
        cursor.execute('INSERT INTO logs (responsavel, acao, detalhes, data_hora) VALUES (?, ?, ?, ?)', 
                       (usuario_responsavel, acao, detalhes, data_hora))
        conn.commit(); conn.close()
    except Exception as e: 
        print(f"Erro no log: {e}")

def get_dias_uteis_passados_mes_atual():
    tz = pytz.timezone('America/Sao_Paulo')
    hoje = datetime.now(tz)
    inicio_mes = hoje.replace(day=1)
    dias_passados = 0
    dia_atual = inicio_mes
    while dia_atual <= hoje:
        if dia_atual.weekday() < 5: 
            dias_passados += 1
        dia_atual += timedelta(days=1)
    return dias_passados

class User(UserMixin):
    def __init__(self, id, username, nome, cargo, role, ultimo_login, ativo):
        self.id = id; self.username = username; self.nome = nome; self.cargo = cargo
        self.role = role; self.ultimo_login = ultimo_login; self.ativo = ativo

@login_manager.user_loader
def load_user(user_id):
    conn = get_db(DB_USERS); cursor = conn.cursor()
    cursor.execute("SELECT * FROM usuarios WHERE id = ?", (user_id,))
    data = cursor.fetchone(); conn.close()
    if data: 
        return User(data['id'], data['username'], data['nome_completo'], data['cargo'], data['role'], data['ultimo_login'], data['ativo'])
    return None

# --- INICIALIZAÇÃO DOS BANCOS DISTINTOS ---
def init_dbs():
    # BANCO DE VENDAS (AGORA COM CORTESIA)
    conn_vendas = get_db(DB_SALES)
    conn_vendas.execute('''CREATE TABLE IF NOT EXISTS resultados (id INTEGER PRIMARY KEY AUTOINCREMENT, consultor TEXT, data_venda TEXT, receita REAL, mig_fixa_bl INTEGER, mig_movel INTEGER, altas_portin INTEGER, bl INTEGER, aparelhos REAL, status TEXT DEFAULT 'Concluido', excluido INTEGER DEFAULT 0, produtos_rs REAL DEFAULT 0.0, cortesia REAL DEFAULT 0.0)''')
    
    # Migração Segura: Adiciona a coluna cortesia se ela não existir no banco atual
    try:
        conn_vendas.execute("ALTER TABLE resultados ADD COLUMN cortesia REAL DEFAULT 0.0")
    except sqlite3.OperationalError:
        pass # A coluna já existe, segue o jogo!
        
    conn_vendas.commit(); conn_vendas.close()

    conn_metas = get_db(DB_METAS)
    conn_metas.execute('''CREATE TABLE IF NOT EXISTS metas (id INTEGER PRIMARY KEY AUTOINCREMENT, nome_categoria TEXT, receita REAL, mig_fixa_bl INTEGER, mig_movel INTEGER, altas_portin INTEGER, bl INTEGER, aparelhos REAL, bonus REAL, mes_referencia TEXT, criado_por TEXT, data_criacao TEXT)''')
    conn_metas.execute('''CREATE TABLE IF NOT EXISTS config_sistema (chave TEXT PRIMARY KEY, valor TEXT)''')
    conn_metas.execute('''CREATE TABLE IF NOT EXISTS metas_semestrais_config (id INTEGER PRIMARY KEY AUTOINCREMENT, semestre TEXT, nivel TEXT, total_receita REAL DEFAULT 0, total_aparelhos REAL DEFAULT 0, total_bl REAL DEFAULT 0, total_dados_avancados REAL DEFAULT 0, total_voz_avancada REAL DEFAULT 0, total_licencas_office REAL DEFAULT 0, total_servicos_tech REAL DEFAULT 0, ultima_atualizacao TEXT)''')
    conn_metas.commit(); conn_metas.close()

    conn_logs = get_db(DB_LOGS)
    conn_logs.execute('''CREATE TABLE IF NOT EXISTS logs (id INTEGER PRIMARY KEY AUTOINCREMENT, responsavel TEXT, acao TEXT, detalhes TEXT, data_hora TEXT)''')
    conn_logs.commit(); conn_logs.close()

    conn_users = get_db(DB_USERS)
    conn_users.execute('''CREATE TABLE IF NOT EXISTS usuarios (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE, password_hash TEXT, nome_completo TEXT, cargo TEXT, role TEXT, ultimo_login TEXT, ativo INTEGER DEFAULT 1, excluido INTEGER DEFAULT 0)''')
    
    # Admin inicial criado a partir de ADMIN_USUARIO / ADMIN_SENHA
    admin_usuario = os.environ.get('ADMIN_USUARIO', 'admin').lower().strip()
    cursor = conn_users.cursor()
    cursor.execute("SELECT * FROM usuarios WHERE username = ?", (admin_usuario,))
    if not cursor.fetchone():
        admin_senha = os.environ.get('ADMIN_SENHA')
        if not admin_senha:
            # Sem ADMIN_SENHA: gera uma senha aleatória e mostra uma única vez
            admin_senha = secrets.token_urlsafe(12)
            print(f"[INFO] Admin '{admin_usuario}' criado com senha gerada: {admin_senha}  (guarde-a, não será exibida novamente)")
        hashed = generate_password_hash(admin_senha)
        cursor.execute("INSERT INTO usuarios (username, password_hash, nome_completo, cargo, role, ultimo_login, ativo, excluido) VALUES (?, ?, ?, ?, ?, ?, 1, 0)", (admin_usuario, hashed, "Administrador", "Administrador", "admin", "Nunca"))
    conn_users.commit(); conn_users.close()

init_dbs()

# ==========================================
# ROTAS DE PÁGINAS (FRONT-END MODULAR)
# ==========================================
@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated: return redirect(url_for('index'))
    if request.method == 'POST':
        username = request.form.get('username').lower().strip()
        password = request.form.get('password')
        conn = get_db(DB_USERS); cursor = conn.cursor()
        cursor.execute("SELECT * FROM usuarios WHERE username = ? AND excluido = 0", (username,))
        user_data = cursor.fetchone()
        
        if user_data and check_password_hash(user_data['password_hash'], password):
            if user_data['ativo'] == 0: 
                flash('Acesso Bloqueado.', 'error'); conn.close(); return render_template('login.html')
            
            agora = datetime.now(pytz.timezone('America/Sao_Paulo')).strftime("%d/%m/%Y %H:%M")
            cursor.execute("UPDATE usuarios SET ultimo_login = ? WHERE id = ?", (agora, user_data['id']))
            conn.commit(); conn.close()
            
            login_user(User(user_data['id'], user_data['username'], user_data['nome_completo'], user_data['cargo'], user_data['role'], agora, user_data['ativo']))
            registrar_log(user_data['nome_completo'], "LOGIN", "Acessou o sistema.")
            return redirect(url_for('index'))
        else: 
            conn.close(); flash('Credenciais inválidas.', 'error')
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout(): 
    registrar_log(current_user.nome, "LOGOUT", "Saiu do sistema.")
    logout_user()
    return redirect(url_for('login'))

@app.route('/')
@login_required
def index(): return render_template('index.html', usuario=current_user)

@app.route('/lancamentos')
@login_required
def lancamentos():
    if current_user.role == 'consultor': return redirect(url_for('index'))
    return render_template('lancamentos.html', usuario=current_user)

@app.route('/metas_tela')
@login_required
def metas_tela():
    if current_user.role == 'consultor': return redirect(url_for('index'))
    return render_template('metas.html', usuario=current_user)

@app.route('/premiacao')
@login_required
def premiacao():
    if current_user.role == 'consultor': return redirect(url_for('index'))
    return render_template('premiacao.html', usuario=current_user)

@app.route('/usuarios_tela')
@login_required
def usuarios_tela():
    if current_user.role == 'consultor': return redirect(url_for('index'))
    return render_template('usuarios.html', usuario=current_user)

@app.route('/backup')
@login_required
def backup_tela():
    if current_user.role != 'admin': return redirect(url_for('index'))
    return render_template('backup.html', usuario=current_user)


# ==========================================
# ROTAS DE API (BACKEND & BASE DE DADOS)
# ==========================================

# --- API: DADOS DE VENDAS ---
@app.route('/api/dados', methods=['GET'])
@login_required
def get_dados():
    conn = get_db(DB_SALES); cursor = conn.cursor()
    mostrar_excluidos = request.args.get('excluidos') == 'true'
    query = "SELECT * FROM resultados WHERE 1=1"
    params = []
    if current_user.role == 'consultor': 
        query += " AND consultor = ? AND excluido = 0"
        params.append(current_user.nome)
    elif not mostrar_excluidos: query += " AND excluido = 0"
    cursor.execute(query, params)
    rows = cursor.fetchall(); conn.close()
    return jsonify([dict(r) for r in rows])

@app.route('/api/salvar', methods=['POST'])
@login_required
def salvar_dados():
    if current_user.role == 'consultor': return jsonify({"error": "Negado"}), 403
    data = request.json
    data_venda = data.get('data_venda', datetime.now().strftime("%Y-%m-%d")) 
    produtos_rs = float(data.get('produtos_rs', 0.0))
    cortesia = float(data.get('cortesia', 0.0)) # NOVO CAMPO CORTESIA
    
    conn = get_db(DB_SALES); cursor = conn.cursor()
    if data.get('id'):
        cursor.execute('''UPDATE resultados SET consultor=?, data_venda=?, receita=?, mig_fixa_bl=?, mig_movel=?, altas_portin=?, bl=?, aparelhos=?, status=?, produtos_rs=?, cortesia=? WHERE id=?''', (data['consultor'], data_venda, data['receita'], data['mig_fixa_bl'], data['mig_movel'], data['altas_portin'], data['bl'], data['aparelhos'], data['status'], produtos_rs, cortesia, data['id']))
        acao = "ATUALIZAR_VENDA"
    else:
        cursor.execute('''INSERT INTO resultados (consultor, data_venda, receita, mig_fixa_bl, mig_movel, altas_portin, bl, aparelhos, status, produtos_rs, cortesia) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''', (data['consultor'], data_venda, data['receita'], data['mig_fixa_bl'], data['mig_movel'], data['altas_portin'], data['bl'], data['aparelhos'], data.get('status','Concluido'), produtos_rs, cortesia))
        acao = "NOVA_VENDA"
    conn.commit(); conn.close()
    registrar_log(current_user.nome, acao, f"Venda de {data['consultor']}")
    return jsonify({"message": "Salvo"}), 201

@app.route('/api/excluir_venda/<int:id>', methods=['DELETE'])
@login_required
def excluir_venda(id):
    conn = get_db(DB_SALES); cursor = conn.cursor()
    cursor.execute("UPDATE resultados SET excluido = 1 WHERE id = ?", (id,))
    conn.commit(); conn.close()
    registrar_log(current_user.nome, "EXCLUSAO_VENDA", f"ID {id} para lixeira.")
    return jsonify({"message": "Excluído"}), 200

# --- API: CÉREBRO DE BI ---
@app.route('/api/bi/performance', methods=['GET'])
@login_required
def api_bi_performance():
    try:
        hoje = datetime.now(pytz.timezone('America/Sao_Paulo'))
        mes_str = hoje.strftime("%Y-%m") 
        
        conn_metas = get_db(DB_METAS); cursor_m = conn_metas.cursor()
        cursor_m.execute("SELECT valor FROM config_sistema WHERE chave = 'dias_uteis'")
        row = cursor_m.fetchone(); dias_totais = int(float(row['valor'])) if row and row['valor'] else 22
        
        semestre_req = f"{1 if hoje.month <= 6 else 2}-{hoje.year}"
        cursor_m.execute("SELECT total_receita FROM metas_semestrais_config WHERE semestre = ? AND nivel = 'GLOBAL'", (semestre_req,))
        meta_row = cursor_m.fetchone()
        meta_mes = (float(meta_row['total_receita'] or 0) / 6) if meta_row and meta_row['total_receita'] else 50000.0 
        conn_metas.close()

        conn_vendas = get_db(DB_SALES); cursor_v = conn_vendas.cursor()
        
        query_vendas = "SELECT SUM(receita) as total FROM resultados WHERE data_venda LIKE ? AND excluido = 0 AND status = 'Concluido'"
        params_vendas = [f"{mes_str}%"]
        
        if current_user.role == 'consultor':
             query_vendas += " AND consultor = ?"
             params_vendas.append(current_user.nome)
             nome_analise = current_user.nome
        else:
             nome_analise = "Equipe"
             
        cursor_v.execute(query_vendas, params_vendas)
        res = cursor_v.fetchone(); realizado_mes = float(res['total'] or 0) if res and res['total'] else 0.0
        conn_vendas.close()

        dias_passados = min(get_dias_uteis_passados_mes_atual(), dias_totais)
        dias_restantes = max(0, dias_totais - dias_passados)
        
        resultado_bi = calculos.analisar_performance_consultor(nome_analise, realizado_mes, meta_mes, dias_passados, dias_restantes)
        resultado_bi['metadados'] = {"dias_uteis_mes": dias_totais, "dias_uteis_passados": dias_passados, "dias_uteis_restantes": dias_restantes}
        return jsonify(resultado_bi)
    except Exception as e:
        print(f"Erro no Cérebro BI: {e}")
        return jsonify({"error": "Erro no servidor BI"}), 500

# --- API: LOGS E USUÁRIOS ---
@app.route('/api/logs', methods=['GET'])
@login_required
def get_logs():
    conn = get_db(DB_LOGS); cursor = conn.cursor()
    cursor.execute("SELECT data_hora as data, responsavel, acao, detalhes FROM logs ORDER BY id DESC LIMIT 500")
    rows = cursor.fetchall(); conn.close()
    return jsonify([dict(r) for r in rows])

@app.route('/api/usuarios', methods=['GET'])
@login_required
def get_usuarios():
    conn = get_db(DB_USERS); cursor = conn.cursor()
    cursor.execute("SELECT id, username, nome_completo as nome, cargo, role, ultimo_login, ativo FROM usuarios WHERE excluido = 0 ORDER BY nome_completo ASC")
    rows = cursor.fetchall(); conn.close()
    return jsonify([dict(r) for r in rows])

@app.route('/api/lista_consultores', methods=['GET'])
@login_required
def get_lista_consultores():
    conn = get_db(DB_USERS); cursor = conn.cursor()
    cursor.execute("SELECT nome_completo FROM usuarios WHERE ativo = 1 AND excluido = 0 ORDER BY nome_completo ASC")
    rows = cursor.fetchall(); conn.close()
    return jsonify([r['nome_completo'] for r in rows])

@app.route('/api/usuarios/salvar', methods=['POST'])
@login_required
def salvar_usuario():
    if current_user.role != 'admin': return jsonify({'error': 'Negado'}), 403
    data = request.json
    
    user_id = data.get('id')
    username = data.get('username').lower().strip()
    nome = data.get('nome')
    cargo = data.get('cargo')
    role = data.get('role')
    password = data.get('password')

    conn = get_db(DB_USERS); cursor = conn.cursor()

    try:
        if user_id:
            if password: 
                hashed = generate_password_hash(password)
                cursor.execute('''UPDATE usuarios SET username=?, password_hash=?, nome_completo=?, cargo=?, role=? WHERE id=?''', (username, hashed, nome, cargo, role, user_id))
            else: 
                cursor.execute('''UPDATE usuarios SET username=?, nome_completo=?, cargo=?, role=? WHERE id=?''', (username, nome, cargo, role, user_id))
            acao = "ATUALIZAR_USUARIO"; msg = f"Editou {nome}"
        else:
            cursor.execute("SELECT id FROM usuarios WHERE username = ?", (username,))
            if cursor.fetchone():
                conn.close(); return jsonify({"error": "Este login já existe!"}), 400
                
            hashed = generate_password_hash(password) if password else generate_password_hash(SENHA_PADRAO)
            cursor.execute("INSERT INTO usuarios (username, password_hash, nome_completo, cargo, role, ultimo_login, ativo, excluido) VALUES (?, ?, ?, ?, ?, ?, 1, 0)", (username, hashed, nome, cargo, role, "Nunca"))
            acao = "CRIAR_USUARIO"; msg = f"Criou {nome}"

        conn.commit(); conn.close()
        registrar_log(current_user.nome, acao, msg)
        return jsonify({"message": "Salvo"}), 200
    except Exception as e: 
        conn.close(); return jsonify({"error": str(e)}), 500

@app.route('/api/usuarios/toggle_status/<int:id>', methods=['POST'])
@login_required
def toggle_status(id):
    if current_user.role != 'admin': return jsonify({'error': 'Negado'}), 403
    conn = get_db(DB_USERS); cursor = conn.cursor()
    cursor.execute("SELECT ativo, nome_completo FROM usuarios WHERE id=?", (id,)); data = cursor.fetchone()
    novo = 0 if data['ativo'] == 1 else 1
    cursor.execute("UPDATE usuarios SET ativo = ? WHERE id = ?", (novo, id))
    conn.commit(); conn.close()
    registrar_log(current_user.nome, "BLOQUEIO", f"Alterou status de {data['nome_completo']}")
    return jsonify({"message": "Ok"}), 200

@app.route('/api/usuarios/deletar/<int:id>', methods=['DELETE'])
@login_required
def deletar_usuario(id):
    if current_user.role != 'admin': return jsonify({'error': 'Negado'}), 403
    if id == current_user.id: return jsonify({'error': 'Erro'}), 400
    conn = get_db(DB_USERS); cursor = conn.cursor()
    cursor.execute("UPDATE usuarios SET excluido = 1, ativo = 0 WHERE id = ?", (id,))
    conn.commit(); conn.close()
    registrar_log(current_user.nome, "EXCLUSAO", f"Excluiu utilizador.")
    return jsonify({"message": "Ok"}), 200

@app.route('/api/admin/resetar_senha/<int:id>', methods=['POST'])
@login_required
def resetar_senha_usuario(id):
    if current_user.role != 'admin': return jsonify({'error': 'Negado'}), 403
    conn = get_db(DB_USERS); cursor = conn.cursor()
    cursor.execute("UPDATE usuarios SET password_hash = ? WHERE id = ?", (generate_password_hash(SENHA_PADRAO), id))
    conn.commit(); conn.close()
    registrar_log(current_user.nome, "RESET_SENHA", "Resetou senha para a senha padrão.")
    return jsonify({"message": "Ok"}), 200

# --- API: METAS ---
@app.route('/api/metas', methods=['GET'])
@login_required
def get_metas():
    mes = request.args.get('mes') or datetime.now(pytz.timezone('America/Sao_Paulo')).strftime("%m/%Y")
    conn = get_db(DB_METAS); cursor = conn.cursor()
    cursor.execute("SELECT * FROM metas WHERE mes_referencia = ? ORDER BY receita ASC", (mes,))
    rows = cursor.fetchall(); conn.close()
    return jsonify([dict(r) for r in rows])

@app.route('/api/metas/salvar', methods=['POST'])
@login_required
def salvar_meta():
    if current_user.role == 'consultor': return jsonify({"error": "Negado"}), 403
    data = request.json; mes_ref = data.get('mes_referencia')
    conn = get_db(DB_METAS); cursor = conn.cursor()
    cursor.execute("SELECT id FROM metas WHERE nome_categoria = ? AND mes_referencia = ?", (data['categoria'], mes_ref))
    existente = cursor.fetchone()
    if existente: cursor.execute('''UPDATE metas SET receita=?, mig_fixa_bl=?, mig_movel=?, altas_portin=?, bl=?, aparelhos=?, bonus=? WHERE id=?''', (data['receita'], data['mig_fixa_bl'], data['mig_movel'], data['altas_portin'], data['bl'], data['aparelhos'], data['bonus'], existente['id']))
    else: cursor.execute('''INSERT INTO metas (nome_categoria, receita, mig_fixa_bl, mig_movel, altas_portin, bl, aparelhos, bonus, mes_referencia, criado_por, data_criacao) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''', (data['categoria'], data['receita'], data['mig_fixa_bl'], data['mig_movel'], data['altas_portin'], data['bl'], data['aparelhos'], data['bonus'], mes_ref, current_user.nome, datetime.now().strftime("%d/%m/%Y")))
    conn.commit(); conn.close(); registrar_log(current_user.nome, "META_MENSAL", f"Configurou {data['categoria']}")
    return jsonify({"message": "Salvo"})

@app.route('/api/semestral/dados', methods=['GET'])
@login_required
def get_semestral_global():
    sem_req = request.args.get('semestre') or f"{1 if datetime.now().month <= 6 else 2}-{datetime.now().year}"
    conn = get_db(DB_METAS); cursor = conn.cursor()
    cursor.execute("SELECT * FROM metas_semestrais_config WHERE semestre = ? AND nivel = 'GLOBAL'", (sem_req,))
    row = cursor.fetchone(); conn.close()
    if row: return jsonify({'semestre': row['semestre'], 'meta_total': row['total_receita'], 'meta_aparelhos': row['total_aparelhos'], 'meta_bl': row['total_bl'], 'meta_dados_avancados': row['total_dados_avancados'], 'meta_voz_avancada': row['total_voz_avancada'], 'meta_licencas_office': row['total_licencas_office'], 'meta_servicos_tech': row['total_servicos_tech']})
    return jsonify({})

@app.route('/api/semestral/salvar', methods=['POST'])
@login_required
def save_semestral_global():
    if not check_admin_access(): return jsonify({'error': 'Restrito'}), 403
    data = request.json; sem = data.get('semestre')
    conn = get_db(DB_METAS); cursor = conn.cursor()
    cursor.execute("SELECT id FROM metas_semestrais_config WHERE semestre = ? AND nivel = 'GLOBAL'", (sem,))
    row = cursor.fetchone()
    params = (float(data.get('meta_receita', 0)), float(data.get('meta_aparelhos', 0)), float(data.get('meta_bl', 0)), float(data.get('meta_dados_avancados', 0)), float(data.get('meta_voz_avancada', 0)), float(data.get('meta_licencas_office', 0)), float(data.get('meta_servicos_tech', 0)))
    if row: cursor.execute("UPDATE metas_semestrais_config SET total_receita=?, total_aparelhos=?, total_bl=?, total_dados_avancados=?, total_voz_avancada=?, total_licencas_office=?, total_servicos_tech=? WHERE id=?", params + (row['id'],))
    else: cursor.execute("INSERT INTO metas_semestrais_config (total_receita, total_aparelhos, total_bl, total_dados_avancados, total_voz_avancada, total_licencas_office, total_servicos_tech, semestre, nivel) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'GLOBAL')", params + (sem,))
    conn.commit(); conn.close(); registrar_log(current_user.nome, "META_SEMESTRAL", f"Atualizou Rumo ao Ouro")
    return jsonify({'ok': True})

# ==========================================
# ROTAS DO COFRE DE BACKUP E IMPORTAÇÃO
# ==========================================
@app.route('/api/backup/db/<tipo>', methods=['GET'])
@login_required
def backup_db_individual(tipo):
    if not check_admin_access(): return jsonify({'error': 'Negado'}), 403
    if tipo not in DB_MAP: return jsonify({'error': 'Cofre inválido'}), 400
    db_file = DB_MAP[tipo][0]
    registrar_log(current_user.nome, "BACKUP DB", f"Transferiu ficheiro raiz de {tipo.upper()}")
    return send_file(db_file, as_attachment=True, download_name=f"Backup_{tipo.capitalize()}_{datetime.now().strftime('%d_%m_%Y')}.db")

@app.route('/api/backup/excel/<tipo>', methods=['GET'])
@login_required
def backup_excel_individual(tipo):
    if not check_admin_access(): return jsonify({'error': 'Negado'}), 403
    if tipo not in DB_MAP: return jsonify({'error': 'Cofre inválido'}), 400
    db_file, tabelas = DB_MAP[tipo]
    wb = Workbook()
    if 'Sheet' in wb.sheetnames: del wb['Sheet']
    
    if os.path.exists(db_file):
        conn = get_db(db_file); cursor = conn.cursor()
        for tab in tabelas:
            ws = wb.create_sheet(tab.capitalize())
            cursor.execute(f"SELECT * FROM {tab}")
            rows = cursor.fetchall()
            cursor.execute(f"PRAGMA table_info({tab})")
            ws.append([c[1] for c in cursor.fetchall()])
            for r in rows: ws.append(list(r))
        conn.close()

    buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
    registrar_log(current_user.nome, "BACKUP EXCEL", f"Transferiu Excel de {tipo.upper()}")
    return send_file(buffer, as_attachment=True, download_name=f"Relatorio_{tipo.capitalize()}_{datetime.now().strftime('%d_%m_%Y')}.xlsx", mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/api/admin/exportar_total')
@login_required
def exportar_total():
    if not check_admin_access(): return jsonify({'error': 'Negado'}), 403
    wb = Workbook()
    if 'Sheet' in wb.sheetnames: del wb['Sheet']
    
    for _, (db_file, tabelas) in DB_MAP.items():
        if os.path.exists(db_file):
            conn = sqlite3.connect(db_file)
            cursor = conn.cursor()
            for tab in tabelas:
                ws = wb.create_sheet(tab.capitalize())
                cursor.execute(f"SELECT * FROM {tab}")
                rows = cursor.fetchall()
                cursor.execute(f"PRAGMA table_info({tab})")
                ws.append([c[1] for c in cursor.fetchall()])
                for r in rows: ws.append(list(r))
            conn.close()

    buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
    registrar_log(current_user.nome, "BACKUP GERAL", "Baixou planilha contendo os 4 cofres")
    return send_file(buffer, as_attachment=True, download_name=f"Backup_Mestre_{datetime.now().strftime('%d_%m_%Y')}.xlsx", mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/api/admin/importar_excel', methods=['POST'])
@login_required
def importar_excel():
    if not check_admin_access(): return jsonify({'error': 'Acesso negado'}), 403
    if 'arquivo' not in request.files: return jsonify({'error': 'Nenhum ficheiro enviado'}), 400

    arquivo = request.files['arquivo']
    if not arquivo.filename.endswith('.xlsx'): return jsonify({'error': 'Formato inválido. Use um ficheiro Excel (.xlsx)'}), 400

    try:
        wb = load_workbook(arquivo, data_only=True)
        registos_importados = 0
        sheet_map = {str(name).lower().strip(): str(name) for name in wb.sheetnames}

        for tipo, (db_file, tabelas) in DB_MAP.items():
            conn = get_db(db_file)
            cursor = conn.cursor()

            for tab in tabelas:
                if tab.lower() in sheet_map:
                    real_sheet_name = sheet_map[tab.lower()]
                    ws = wb[real_sheet_name]
                    rows = list(ws.rows)
                    if len(rows) > 1: 
                        excel_headers = [str(c.value).strip() if c.value else "" for c in rows[0]]
                        cursor.execute(f"PRAGMA table_info({tab})")
                        db_columns = [col['name'] for col in cursor.fetchall()]
                        
                        colunas_validas = []
                        indices_validos = []
                        for i, col_name in enumerate(excel_headers):
                            if col_name in db_columns:
                                colunas_validas.append(col_name)
                                indices_validos.append(i)
                        
                        if not colunas_validas: continue 
                        
                        cursor.execute(f"DELETE FROM {tab}")
                        placeholders = ",".join(["?"] * len(colunas_validas))
                        col_str = ",".join(colunas_validas)
                        sql = f"INSERT INTO {tab} ({col_str}) VALUES ({placeholders})"

                        for row in rows[1:]:
                            valores = []
                            linha_vazia = True
                            for i in indices_validos:
                                val = row[i].value
                                if isinstance(val, datetime): val = val.strftime('%Y-%m-%d')
                                elif str(val).strip() in ['NaN', '#N/A', 'None']: val = 0.0 if 'receita' in colunas_validas else ""
                                if val is not None and str(val).strip() != "": linha_vazia = False
                                valores.append(val)
                            
                            if not linha_vazia:
                                try:
                                    cursor.execute(sql, valores)
                                    registos_importados += 1
                                except Exception as erro_linha: pass
            
            conn.commit(); conn.close()
        registrar_log(current_user.nome, "RESTAURAÇÃO", f"Sistema restaurado via Excel ({registos_importados} linhas).")
        return jsonify({'message': f'Sucesso! {registos_importados} registos transferidos.'}), 200
    except Exception as e: return jsonify({'error': f'Estrutura incompatível: {str(e)}'}), 500

if __name__ == '__main__':
    app.run(debug=os.environ.get('FLASK_DEBUG') == '1', host='0.0.0.0', port=int(os.environ.get('PORT', 5000)))